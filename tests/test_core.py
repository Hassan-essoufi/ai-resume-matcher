import io

import pytest
from docx import Document

from matcher.config import Settings
from matcher.extraction import RuleBasedExtractor
from matcher.ingestion import UnsupportedFileError, extract_text
from matcher.pii import regex_scrub
from matcher.schemas import CandidateProfile, JobProfile
from matcher.scoring import experience_fit, match_skills, rescale, score
from matcher.skills import find_skills, normalize_skill, normalize_skills


def test_normalize_aliases():
    assert normalize_skill("K8s") == "kubernetes"
    assert normalize_skill(" Sklearn ") == "scikit-learn"
    assert normalize_skills(["Python", "python3", "PYTHON"]) == ["python"]


def test_find_skills_avoids_false_positives():
    found = find_skills("I rest on weekends and go hiking. Skills: Python, Docker, Postgres")
    assert {"python", "docker", "postgresql"} <= set(found)
    assert "go" not in found and "rest api" not in found


def test_pii_scrub_removes_contacts_keeps_dates_and_skills():
    text = "Karim Tazi\nML Engineer\nkarim@example.com | +212 600 000 001 | github.com/karim\n2019 - 2023 Flask"
    out = regex_scrub(text)
    assert "Karim" not in out and "example.com" not in out and "600 000" not in out
    assert "<NAME>" in out and "<EMAIL>" in out and "<PHONE>" in out
    assert "2019 - 2023" in out and "Flask" in out


def test_scoring_components():
    s = Settings()
    job = JobProfile(title="ML", required_skills=["python", "pytorch"], nice_to_have_skills=["aws"],
                     min_years_experience=4, seniority="mid")
    cand = CandidateProfile(skills=["python", "aws"], years_experience=2, seniority="mid")
    b = score(job, cand, cosine=s.semantic_ceiling, settings=s)
    assert b.matched_required == ["python"] and b.missing_required == ["pytorch"]
    assert b.required_coverage == 0.5 and b.nice_coverage == 1.0 and b.experience_fit == 0.5
    expected = 100 * (s.w_semantic * 1 + s.w_required * 0.5 + s.w_nice * 1 + s.w_experience * 0.5)
    assert abs(b.total - round(expected, 1)) < 1e-6


def test_perfect_candidate_scores_100():
    s = Settings()
    job = JobProfile(title="x", required_skills=["python"], min_years_experience=2)
    cand = CandidateProfile(skills=["python"], years_experience=5, seniority="senior")
    assert score(job, cand, cosine=0.99, settings=s).total == 100.0


def test_helpers():
    assert rescale(0.0, 0.15, 0.75) == 0.0 and rescale(1.0, 0.15, 0.75) == 1.0
    assert experience_fit(3, 0) == 1.0 and experience_fit(1, 4) == 0.25
    assert match_skills(["a", "b"], ["b"]) == (["b"], ["a"])


def test_rule_based_job_extraction_splits_required_and_nice():
    job = RuleBasedExtractor().job("ML Engineer\nRequirements: 3+ years, Python, PyTorch\nNice to have: AWS")
    assert {"python", "pytorch"} <= set(job.required_skills)
    assert job.nice_to_have_skills == ["aws"] and job.min_years_experience == 3


def test_docx_and_txt_ingestion(tmp_path):
    doc = Document()
    doc.add_paragraph("Backend developer with 5 years of Python and FastAPI experience.")
    doc.add_paragraph("Skills: PostgreSQL, Docker, Redis, Kubernetes, CI/CD.")
    buf = io.BytesIO()
    doc.save(buf)
    assert "FastAPI" in extract_text("cv.docx", buf.getvalue())
    assert "Docker" in extract_text("cv.txt", b"Skills: Docker, Python, Kubernetes and a long enough text here.")
    with pytest.raises(UnsupportedFileError):
        extract_text("photo.png", b"...")
