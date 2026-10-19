import io
import re
import unicodedata
from pathlib import Path
from docx import Document
from pypdf import PdfReader

class UnsupportedFileError(ValueError):
    pass


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("•", "-").replace("", "-")      
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)                   
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        raw = _pdf_text(data)
    elif ext == ".docx":
        raw = _docx_text(data)
    elif ext in {".txt", ".md"}:
        raw = data.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileError(f"Unsupported file type: {ext} (use PDF, DOCX or TXT)")
    text = clean_text(raw)
    if len(text) < 50:
        raise UnsupportedFileError(
            f"{filename}: almost no text extracted. Scanned PDFs need OCR first.")
    return text
